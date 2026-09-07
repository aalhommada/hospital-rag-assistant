"""
Turning files into text we can chunk.

A loader has one job: read a file and return `LoadedDocument` — clean text plus
the metadata every later stage needs. Everything downstream is file-format
blind, which is why adding a new format later means adding one function here
and nothing else.

The interesting part is what we keep. Most naive loaders return a wall of text
and throw the structure away. We keep heading levels, because headings are how
a hospital document is organised and they end up being the single most useful
signal both for chunking and for citations.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

SUPPORTED_SUFFIXES = {".md", ".markdown", ".pdf", ".docx"}


@dataclass
class Block:
    """One heading or one paragraph, in the order it appeared in the file."""

    kind: str  # "heading" or "text"
    text: str
    level: int = 0  # heading depth: 1 for "#", 2 for "##", 0 for body text


@dataclass
class LoadedDocument:
    title: str
    source_path: str
    source_type: str
    checksum: str
    blocks: list[Block] = field(default_factory=list)
    category: str = "department"
    reviewed_on: date | None = None


class UnsupportedFileType(Exception):
    pass


def checksum_of(path: Path) -> str:
    """SHA-256 of the raw bytes, so we can skip files that have not changed."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(65536), b""):
            digest.update(block)
    return digest.hexdigest()


def load(path: Path) -> LoadedDocument:
    """Load any supported file. Dispatches on the file extension."""
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown"}:
        return load_markdown(path)
    if suffix == ".pdf":
        return load_pdf(path)
    if suffix == ".docx":
        return load_docx(path)
    raise UnsupportedFileType(f"{path.name}: only {', '.join(sorted(SUPPORTED_SUFFIXES))} are supported")


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

# Optional front matter at the top of a hospital document:
#
#   ---
#   category: preparation
#   reviewed_on: 2026-03-01
#   ---
#
# Both keys are optional. `category` groups documents in the admin and lets
# retrieval be filtered later; `reviewed_on` is shown next to every citation.
FRONT_MATTER = re.compile(r"\A---\s*\n(.*?)\n---\s*\n", re.DOTALL)
HEADING = re.compile(r"^(#{1,6})\s+(.*)$")


def load_markdown(path: Path) -> LoadedDocument:
    raw = path.read_text(encoding="utf-8")
    metadata, body = _split_front_matter(raw)

    blocks: list[Block] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            text = " ".join(line.strip() for line in paragraph).strip()
            if text:
                blocks.append(Block(kind="text", text=text))
            paragraph.clear()

    for line in body.splitlines():
        heading = HEADING.match(line)
        if heading:
            flush()
            blocks.append(Block(kind="heading", text=heading.group(2).strip(), level=len(heading.group(1))))
        elif line.strip():
            paragraph.append(line)
        else:
            # A blank line ends a paragraph. List items stay inside the
            # paragraph they belong to, which keeps "bring these four things"
            # together as one retrievable unit instead of four fragments.
            flush()
    flush()

    title = _first_heading(blocks) or path.stem.replace("-", " ").title()
    return LoadedDocument(
        title=title,
        source_path=str(path),
        source_type="markdown",
        checksum=checksum_of(path),
        blocks=blocks,
        category=metadata.get("category", "department"),
        reviewed_on=_parse_date(metadata.get("reviewed_on")),
    )


def _split_front_matter(raw: str) -> tuple[dict[str, str], str]:
    match = FRONT_MATTER.match(raw)
    if not match:
        return {}, raw
    metadata = {}
    for line in match.group(1).splitlines():
        if ":" in line:
            key, _, value = line.partition(":")
            metadata[key.strip()] = value.strip()
    return metadata, raw[match.end():]


def _parse_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError:
        return None


def _first_heading(blocks: list[Block]) -> str | None:
    return next((b.text for b in blocks if b.kind == "heading"), None)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def load_pdf(path: Path) -> LoadedDocument:
    """
    Extract text from a PDF, treating short standalone lines as headings.

    A PDF has no heading tags — only glyphs at positions. The heuristic below
    (a short line, no closing punctuation, followed by a blank line) recovers
    most headings in a plainly formatted document, which is what hospital
    leaflets usually are. It will not survive a two-column research paper, and
    it cannot read a scan at all; both of those need a layout-aware parser.
    That is a real limit of this project, and the README says so.
    """
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[Block] = []

    for page in reader.pages:
        text = page.extract_text() or ""
        for raw_paragraph in re.split(r"\n\s*\n", text):
            paragraph = " ".join(raw_paragraph.split())
            if not paragraph:
                continue
            if _looks_like_heading(paragraph):
                blocks.append(Block(kind="heading", text=paragraph, level=2))
            else:
                blocks.append(Block(kind="text", text=paragraph))

    metadata_title = (reader.metadata.title if reader.metadata else None) or ""
    title = metadata_title.strip() or _first_heading(blocks) or path.stem.replace("-", " ").title()
    return LoadedDocument(
        title=title,
        source_path=str(path),
        source_type="pdf",
        checksum=checksum_of(path),
        blocks=blocks,
    )


def _looks_like_heading(paragraph: str) -> bool:
    return (
        len(paragraph) <= 70
        and not paragraph.endswith((".", ",", ";", ":"))
        and len(paragraph.split()) <= 10
    )


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

def load_docx(path: Path) -> LoadedDocument:
    """Word keeps real heading styles, so this loader is the most reliable one."""
    import docx

    document = docx.Document(str(path))
    blocks: list[Block] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading"):
            level = int(style.replace("heading", "").strip() or 1)
            blocks.append(Block(kind="heading", text=text, level=level))
        elif style == "title":
            blocks.append(Block(kind="heading", text=text, level=1))
        else:
            blocks.append(Block(kind="text", text=text))

    # Tables are flattened row by row. A row read as "Department | Floor | Phone"
    # keeps its columns together, which is enough for the simple tables in
    # hospital documents.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(Block(kind="text", text=" | ".join(cells)))

    title = _first_heading(blocks) or path.stem.replace("-", " ").title()
    return LoadedDocument(
        title=title,
        source_path=str(path),
        source_type="docx",
        checksum=checksum_of(path),
        blocks=blocks,
    )
